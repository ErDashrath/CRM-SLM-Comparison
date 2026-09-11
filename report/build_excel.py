"""
Phase 5 -- assemble results/eval_results.json + the research agent's
tradeoffs writeup into a single Excel workbook, three sheets:

  Per-Query Results  -- one row per (query, variant), all raw scores
  Aggregate Summary   -- per-variant roll-up, matches research_agent.py's numbers
  Tradeoffs            -- the research agent's writeup, not hand-authored

Usage:
    python -m report.build_excel
"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent
RESULTS_PATH = PROJECT_ROOT / "results" / "eval_results.json"
BUNDLE_PATH = PROJECT_ROOT / "results" / "tradeoffs_bundle.json"
KD_MANIFEST_PATH = PROJECT_ROOT / "data" / "kd_manifest.json"
SFT_MANIFEST_PATH = PROJECT_ROOT / "data" / "sft_manifest.json"


def build_per_query_df(results: list[dict]) -> pd.DataFrame:
    rows = []
    for r in results:
        rows.append(
            {
                "variant": r["variant"],
                "account_id": r["account_id"] or "portfolio",
                "query": r["query"],
                "parsed_ok": r["parsed_ok"],
                "guardrail_pass": r["guardrail_pass"],
                "guardrail_violations": "; ".join(r["guardrail_violations"]) if r["guardrail_violations"] else "",
                "judge_correctness": r["judge"]["correctness"],
                "judge_completeness": r["judge"]["completeness"],
                "judge_risk_surfacing": r["judge"]["risk_surfacing"],
                "judge_rationale": r["judge"]["rationale"],
                "output_tokens": r["output_tokens"],
                "tokens_per_second": r["tokens_per_second"],
                "elapsed_seconds": r["elapsed_seconds"],
                "response_text": r["response_text"],
            }
        )
    return pd.DataFrame(rows)


def build_aggregate_df(aggregate_summary: dict) -> pd.DataFrame:
    kd_manifest = json.loads(KD_MANIFEST_PATH.read_text()) if KD_MANIFEST_PATH.exists() else {}
    sft_manifest = json.loads(SFT_MANIFEST_PATH.read_text()) if SFT_MANIFEST_PATH.exists() else {}

    training_data = {
        "base": {"training_examples": 0, "human_review_minutes": 0, "selection_method": "n/a -- no fine-tuning"},
        "kd": {
            "training_examples": kd_manifest.get("accepted", "?"),
            "human_review_minutes": kd_manifest.get("human_review_minutes", 0),
            "selection_method": "teacher-generated, guardrail-filtered only",
        },
        "sft": {
            "training_examples": sft_manifest.get("total_selected", "?"),
            "human_review_minutes": sft_manifest.get("human_review_minutes", 0),
            "selection_method": sft_manifest.get("selection_method", "?"),
        },
    }

    rows = []
    for variant in ("base", "kd", "sft"):
        stats = aggregate_summary.get(variant, {})
        train = training_data[variant]
        rows.append(
            {
                "variant": variant,
                "training_examples": train["training_examples"],
                "human_review_minutes": train["human_review_minutes"],
                "selection_method": train["selection_method"],
                "eval_queries": stats.get("n_queries"),
                "parsed_ok": stats.get("parsed_ok"),
                "guardrail_pass": stats.get("guardrail_pass"),
                "avg_judge_correctness": stats.get("avg_correctness"),
                "avg_judge_completeness": stats.get("avg_completeness"),
                "avg_judge_risk_surfacing": stats.get("avg_risk_surfacing"),
                "avg_output_tokens": stats.get("avg_output_tokens"),
                "avg_tokens_per_second": stats.get("avg_tokens_per_second"),
            }
        )
    return pd.DataFrame(rows)


def _autofit_and_wrap(ws, wrap_columns: set[str] = frozenset()) -> None:
    header = [cell.value for cell in ws[1]]
    for col_idx, col_name in enumerate(header, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 60 if col_name in wrap_columns else min(
            max(len(str(col_name)), 12), 30
        )
        if col_name in wrap_columns:
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)


def main() -> None:
    results = json.loads(RESULTS_PATH.read_text())
    bundle = json.loads(BUNDLE_PATH.read_text())

    per_query_df = build_per_query_df(results)
    aggregate_df = build_aggregate_df(bundle["aggregate_summary"])

    timestamp = datetime.now().strftime("%Y%m%d-%H%M")
    out_path = PROJECT_ROOT / "results" / f"results-{timestamp}.xlsx"

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        aggregate_df.to_excel(writer, sheet_name="Aggregate Summary", index=False)
        per_query_df.to_excel(writer, sheet_name="Per-Query Results", index=False)

        tradeoffs_ws = writer.book.create_sheet("Tradeoffs")
        tradeoffs_ws["A1"] = "CRM Small-Model Comparison -- Tradeoffs Report"
        tradeoffs_ws["A1"].font = Font(bold=True, size=14)
        tradeoffs_ws["A2"] = f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} via report/research_agent.py"
        tradeoffs_ws["A2"].font = Font(italic=True, size=9)

        row = 4
        for line in bundle["report_text"].split("\n"):
            cell = tradeoffs_ws.cell(row=row, column=1, value=line)
            if line.startswith("## "):
                cell.font = Font(bold=True, size=12)
                cell.value = line[3:]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        tradeoffs_ws.column_dimensions["A"].width = 140

    # openpyxl formatting pass (autofit-ish widths, wrap long text columns)
    from openpyxl import load_workbook

    wb = load_workbook(out_path)
    _autofit_and_wrap(wb["Aggregate Summary"])
    _autofit_and_wrap(
        wb["Per-Query Results"],
        wrap_columns={"query", "guardrail_violations", "judge_rationale", "response_text"},
    )
    wb.save(out_path)

    print(f"Wrote {out_path}")
    print(f"  Aggregate Summary: {len(aggregate_df)} rows")
    print(f"  Per-Query Results: {len(per_query_df)} rows")
    print(f"  Tradeoffs: {len(bundle['report_text'].split(chr(10)))} lines from {bundle['report_path']}")


if __name__ == "__main__":
    main()
